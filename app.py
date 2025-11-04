import streamlit as st
import google.generativeai as genai
from gtts import gTTS
import PyPDF2
import docx
from fpdf import FPDF
import io
import os
from dotenv import load_dotenv
from datetime import datetime

# Load environment variables
load_dotenv()

# Configure Gemini API
genai.configure(api_key=os.getenv('GEMINI_API_KEY'))

# Initialize session state for user management
if 'users' not in st.session_state:
    st.session_state.users = {'admin': 'admin123'}  # Default user
if 'logged_in' not in st.session_state:
    st.session_state.logged_in = False
if 'username' not in st.session_state:
    st.session_state.username = None

# User Authentication Functions
def login_user(username, password):
    if username in st.session_state.users and st.session_state.users[username] == password:
        st.session_state.logged_in = True
        st.session_state.username = username
        return True
    return False

def register_user(username, password):
    if username not in st.session_state.users:
        st.session_state.users[username] = password
        return True
    return False

def logout_user():
    st.session_state.logged_in = False
    st.session_state.username = None

# File Processing Functions
def extract_text_from_pdf(pdf_file):
    pdf_reader = PyPDF2.PdfReader(pdf_file)
    text = ""
    for page in pdf_reader.pages:
        text += page.extract_text()
    return text

def extract_text_from_docx(docx_file):
    doc = docx.Document(docx_file)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text

# Gemini API Functions
def process_text_with_gemini(text, operation, language="English", tone="Neutral", adaptation="General"):
    model = genai.GenerativeModel('gemini-2.0-flash-exp')
    
    # Construct prompt based on operation
    if operation == "summarize":
        prompt = f"""Summarize the following text in {language}.
Tone: {tone}
Style: {adaptation}
Keep the summary concise and capture the key points.

Text: {text}

Summary:"""
    elif operation == "rewrite":
        prompt = f"""Rewrite the following text in {language}.
Tone: {tone}
Style: {adaptation}
Maintain the core meaning while improving clarity and flow.

Text: {text}

Rewritten text:"""
    else:
        prompt = text
    
    try:
        response = model.generate_content(prompt)
        return response.text
    except Exception as e:
        return f"Error: {str(e)}"

# Text-to-Speech Function
def text_to_speech(text, language='en'):
    try:
        # Language mapping for gTTS
        lang_map = {
            'English': 'en',
            'Spanish': 'es',
            'French': 'fr',
            'German': 'de',
            'Hindi': 'hi',
            'Chinese': 'zh-CN',
            'Japanese': 'ja',
            'Korean': 'ko'
        }
        
        lang_code = lang_map.get(language, 'en')
        tts = gTTS(text=text, lang=lang_code, slow=False)
        
        # Save to bytes
        fp = io.BytesIO()
        tts.write_to_fp(fp)
        fp.seek(0)
        return fp
    except Exception as e:
        st.error(f"Error generating audio: {str(e)}")
        return None

# PDF Generation Function
def generate_pdf(text, title="Document"):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, title, ln=True, align='C')
    pdf.ln(10)
    pdf.set_font("Arial", size=12)
    
    # Handle text encoding
    text = text.encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(0, 10, text)
    
    # Return PDF as bytes
    pdf_output = pdf.output(dest='S').encode('latin-1')
    return pdf_output

# Main App
def main():
    st.set_page_config(page_title="AI Text Processor", page_icon="📝", layout="wide")
    
    # Custom CSS
    st.markdown("""
        <style>
        .main-header {
            font-size: 3rem;
            color: #1E88E5;
            text-align: center;
            margin-bottom: 2rem;
        }
        .stButton>button {
            background-color: #1E88E5;
            color: white;
            border-radius: 5px;
        }
        .output-section {
            background-color: #f0f2f6;
            padding: 20px;
            border-radius: 10px;
            margin-top: 20px;
        }
        </style>
    """, unsafe_allow_html=True)
    
    # Login/Register Page
    if not st.session_state.logged_in:
        st.markdown("<h1 class='main-header'>📝 AI Text Processor</h1>", unsafe_allow_html=True)
        
        tab1, tab2 = st.tabs(["Login", "Register"])
        
        with tab1:
            st.subheader("Login")
            username = st.text_input("Username", key="login_username")
            password = st.text_input("Password", type="password", key="login_password")
            
            if st.button("Login"):
                if login_user(username, password):
                    st.success("Logged in successfully!")
                    st.rerun()
                else:
                    st.error("Invalid username or password")
        
        with tab2:
            st.subheader("Register")
            new_username = st.text_input("Username", key="register_username")
            new_password = st.text_input("Password", type="password", key="register_password")
            confirm_password = st.text_input("Confirm Password", type="password")
            
            if st.button("Register"):
                if new_password != confirm_password:
                    st.error("Passwords do not match")
                elif len(new_password) < 6:
                    st.error("Password must be at least 6 characters")
                elif register_user(new_username, new_password):
                    st.success("Registration successful! Please login.")
                else:
                    st.error("Username already exists")
        
        return
    
    # Main Application (After Login)
    col1, col2 = st.columns([6, 1])
    with col1:
        st.markdown("<h1 class='main-header'>📝 AI Text Processor</h1>", unsafe_allow_html=True)
    with col2:
        if st.button("Logout"):
            logout_user()
            st.rerun()
    
    st.write(f"Welcome, **{st.session_state.username}**!")
    
    # Sidebar Configuration
    st.sidebar.header("⚙️ Configuration")
    
    operation = st.sidebar.selectbox(
        "Operation",
        ["Summarize", "Rewrite"]
    )
    
    language = st.sidebar.selectbox(
        "Output Language",
        ["English", "Spanish", "French", "German", "Hindi", "Chinese", "Japanese", "Korean"]
    )
    
    tone = st.sidebar.selectbox(
        "Tone",
        ["Neutral", "Formal", "Casual", "Professional", "Friendly", "Persuasive", "Informative"]
    )
    
    adaptation = st.sidebar.selectbox(
        "Adaptation Style",
        ["General", "Academic", "News Article", "Blog Post", "Technical", "Creative"]
    )
    
    # Main Content Area
    st.header("📄 Input")
    
    input_method = st.radio(
        "Choose input method:",
        ["Type/Paste Text", "Upload File"]
    )
    
    input_text = ""
    
    if input_method == "Type/Paste Text":
        input_text = st.text_area("Enter your text here:", height=200)
    
    elif input_method == "Upload File":
        uploaded_file = st.file_uploader("Upload a file", type=['pdf', 'docx', 'txt'])
        
        if uploaded_file:
            file_type = uploaded_file.name.split('.')[-1].lower()
            
            try:
                if file_type == 'pdf':
                    input_text = extract_text_from_pdf(uploaded_file)
                elif file_type == 'docx':
                    input_text = extract_text_from_docx(uploaded_file)
                elif file_type == 'txt':
                    input_text = uploaded_file.read().decode('utf-8')
                
                st.success(f"✅ File uploaded successfully! Extracted {len(input_text)} characters.")
                with st.expander("📄 View Extracted Text"):
                    st.text_area("Extracted Text:", input_text, height=200)
            except Exception as e:
                st.error(f"Error reading file: {str(e)}")
    
    # Process Button
    if st.button("🚀 Process Text", type="primary"):
        if not input_text:
            st.warning("⚠️ Please provide some input text!")
        else:
            with st.spinner(f"{operation}ing text with AI..."):
                operation_type = "summarize" if operation == "Summarize" else "rewrite"
                output_text = process_text_with_gemini(
                    input_text, 
                    operation_type, 
                    language, 
                    tone, 
                    adaptation
                )
                
                st.session_state.output_text = output_text
                st.session_state.output_language = language
                st.session_state.operation = operation
    
    # Output Section
    if 'output_text' in st.session_state:
        st.markdown("---")
        st.header("✨ Output")
        
        # Display output in a nice container
        st.markdown("<div class='output-section'>", unsafe_allow_html=True)
        st.markdown(f"**{st.session_state.operation}d Text:**")
        st.write(st.session_state.output_text)
        st.markdown("</div>", unsafe_allow_html=True)
        
        # Action Buttons Section
        st.subheader("📥 Actions")
        
        col1, col2, col3 = st.columns(3)
        
        with col1:
            # Download as PDF
            st.markdown("**Download PDF**")
            pdf_bytes = generate_pdf(
                st.session_state.output_text,
                f"{st.session_state.operation} - {datetime.now().strftime('%Y-%m-%d %H:%M')}"
            )
            st.download_button(
                label="📄 Download as PDF",
                data=pdf_bytes,
                file_name=f"{st.session_state.operation.lower()}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf",
                mime="application/pdf",
                use_container_width=True
            )
        
        with col2:
            # Generate and Download Audio
            st.markdown("**Download Audio**")
            if st.button("🎵 Generate Audio File", use_container_width=True):
                with st.spinner("🎧 Generating audio..."):
                    audio_fp = text_to_speech(
                        st.session_state.output_text, 
                        st.session_state.output_language
                    )
                    if audio_fp:
                        st.session_state.audio_data = audio_fp.read()
                        audio_fp.seek(0)
                        st.success("✅ Audio generated successfully!")
        
        with col3:
            # Copy Text
            st.markdown("**Copy Text**")
            st.text_area(
                "Copy from here:", 
                st.session_state.output_text, 
                height=100, 
                key="copy_output",
                label_visibility="collapsed"
            )
        
        # Audio Player and Download Section
        if 'audio_data' in st.session_state:
            st.markdown("---")
            st.subheader("🔊 Audio Output")
            
            col_audio1, col_audio2 = st.columns([2, 1])
            
            with col_audio1:
                st.markdown("**Play Audio:**")
                st.audio(st.session_state.audio_data, format='audio/mp3')
            
            with col_audio2:
                st.markdown("**Download Audio File:**")
                st.download_button(
                    label="⬇️ Download MP3",
                    data=st.session_state.audio_data,
                    file_name=f"{st.session_state.operation.lower()}_audio_{datetime.now().strftime('%Y%m%d_%H%M%S')}.mp3",
                    mime="audio/mp3",
                    use_container_width=True
                )
        
        # Statistics Section
        st.markdown("---")
        st.subheader("📊 Statistics")
        
        stat_col1, stat_col2, stat_col3, stat_col4 = st.columns(4)
        
        with stat_col1:
            st.metric("Word Count", len(st.session_state.output_text.split()))
        
        with stat_col2:
            st.metric("Character Count", len(st.session_state.output_text))
        
        with stat_col3:
            st.metric("Language", st.session_state.output_language)
        
        with stat_col4:
            st.metric("Operation", st.session_state.operation)

if __name__ == "__main__":
    main()